"""Tests for the formatting boundary; exports are validated by running the Jobset."""

import csv
from datetime import datetime
import json
import os
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from pypdf import PdfReader, PdfWriter

import finalize_outputs as formatter
import run_finalize as launcher


class FormattingTests(unittest.TestCase):
    def test_launcher_selects_interpreter_and_preserves_arguments(self):
        args = ["--project-root", "/project with spaces", "--kind", "prepare-assembly"]
        windows = launcher.command(args, "nt", r"C:\Program Files\KiCad\10.0\bin\python.exe")
        linux = launcher.command(args, "posix", "/app/venv/bin/python")
        self.assertEqual(windows[:2], ["py", "-3.14"])
        self.assertEqual(linux[0], "/app/venv/bin/python")
        self.assertEqual(windows[-len(args):], args)
        self.assertEqual(linux[-len(args):], args)
        self.assertEqual(Path(windows[2]).name, "finalize_outputs.py")
        self.assertEqual(Path(linux[1]).name, "finalize_outputs.py")

    def test_launcher_propagates_child_exit_code(self):
        with (patch.object(launcher.subprocess, "run") as run,
              patch.object(launcher.sys, "argv", ["run_finalize.py", "--kind", "assembly"])):
            run.return_value.returncode = 17
            self.assertEqual(launcher.main(), 17)
            run.assert_called_once_with(
                launcher.command(["--kind", "assembly"], launcher.os.name,
                                 launcher.sys.executable),
                check=False,
            )

    @staticmethod
    def _u3d_material(name, diffuse):
        encoded = name.encode()
        payload = (struct.pack("<H", len(encoded)) + encoded + struct.pack("<I", 0x36)
                   + struct.pack("<3f", 0, 0, 0) + struct.pack("<3f", *diffuse)
                   + struct.pack("<3f", 0.2, 0.2, 0.2) + struct.pack("<3f", 0, 0, 0)
                   + struct.pack("<2f", 0.32, 1.0))
        return (struct.pack("<III", 0xFFFFFF54, len(payload), 0) + payload
                + bytes((-len(payload)) % 4))

    def test_recolors_only_named_u3d_pad_material(self):
        copper = (0.7, 0.61, 0.0)
        original = (self._u3d_material("m_Copper_0", copper)
                    + self._u3d_material("m_Pads_0", (0.5, 0.5, 0.5))
                    + self._u3d_material("m_Component_0", (0.5, 0.5, 0.5)))
        changed = formatter.recolor_u3d_pads(original)
        self.assertNotEqual(changed, original)
        self.assertEqual(changed.count(struct.pack("<3f", *copper)), 2)
        self.assertEqual(changed.count(struct.pack("<3f", 0.5, 0.5, 0.5)), 1)
        self.assertEqual(formatter.recolor_u3d_pads(changed), changed)

    def test_recolor_rejects_unexpected_u3d_materials(self):
        data = self._u3d_material("m_Copper_0", (0.7, 0.61, 0.0))
        with self.assertRaisesRegex(ValueError, "m_Pads_0"):
            formatter.recolor_u3d_pads(data)

    def test_requires_jobset_context(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaisesRegex(RuntimeError, "only through"):
                formatter.finalize("assembly", Path.cwd())

    def test_assembly_page_count_follows_enabled_bottom_job(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Outputs.kicad_jobset"
            base = {
                "jobs": [
                    {"id": "prepare", "settings": {"command": "tool --kind prepare-assembly"}},
                    {"id": "bottom", "settings": {"output_filename": "_work/assembly-bottom.pdf"}},
                ],
                "outputs": [{"only": ["prepare", "bottom"]}],
            }
            path.write_text(json.dumps(base))
            self.assertEqual(formatter.assembly_page_count(path), 2)
            base["outputs"][0]["only"].remove("bottom")
            path.write_text(json.dumps(base))
            self.assertEqual(formatter.assembly_page_count(path), 1)

    def test_assembly_variant_uses_no_variant_for_base_selection(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Outputs.kicad_jobset"
            base = {
                "jobs": [
                    {"id": "prepare", "settings": {"command": "tool --kind prepare-assembly"}},
                    {"id": "schematic", "settings": {"variant_name": ""}},
                    {"id": "drawing", "settings": {"variant": ""}},
                ],
                "outputs": [{"only": ["prepare", "schematic", "drawing"]}],
            }
            path.write_text(json.dumps(base))
            self.assertEqual(formatter.assembly_variant_name(path), "No Variant")
            base["jobs"][1]["settings"]["variant_name"] = "Default"
            base["jobs"][2]["settings"]["variant"] = "Default"
            path.write_text(json.dumps(base))
            self.assertEqual(formatter.assembly_variant_name(path), "Default")

    def test_rejects_path_characters(self):
        for value in ("../part", "part/name", "part:name", "", "part."):
            with self.subTest(value=value), self.assertRaises(ValueError):
                formatter.filename_value(value)

    def test_missing_export_does_not_publish_partial_package(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            project = root / "project"
            project.mkdir()
            (project / "test.kicad_pro").write_text(json.dumps({"text_variables": {
                "ProjectTitle": "Test", "ProjectPCBRevision": "1.0",
            }}))
            output = root / "output"
            (output / "_work").mkdir(parents=True)
            (output / "_work/fabrication.pdf").write_bytes(b"native output")
            with patch.dict(os.environ, JOBSET_OUTPUT_WORK_PATH=str(output)):
                with self.assertRaisesRegex(RuntimeError, "envelope.step"):
                    formatter.finalize("fabrication", project)
            self.assertEqual(list(output.iterdir()), [output / "_work"])
            self.assertTrue((output / "_work/fabrication.pdf").exists())

    def test_template_metadata_types_and_literal_text(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "bom.csv"
            target = Path(directory) / "bom.xlsx"
            with source.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream)
                writer.writerow(formatter.HEADERS)
                writer.writerow([1, 2, "R1,R2", "=literal description", "Maker", "00123"])
            published = datetime(2026, 9, 18, 17, 5)
            formatter.create_bom(source, target,
                                 {"ProjectTitle": "Board", "ProjectPCBARevision": "1.0"},
                                 "No Variant", published)
            sheet = openpyxl.load_workbook(target).active
            self.assertEqual(sheet["C4"].value, "1.0")
            self.assertEqual(sheet["C4"].number_format, "@")
            self.assertEqual(sheet["C5"].value, published)
            self.assertEqual(sheet["C6"].value, "No Variant")
            self.assertEqual(sheet["A9"].value, 1)
            self.assertEqual(sheet["B9"].value, 2)
            self.assertEqual(sheet["D9"].data_type, "s")
            self.assertEqual(sheet["F9"].value, "00123")
            self.assertIn("A1:C1", str(sheet.merged_cells))
            template = openpyxl.load_workbook(Path(formatter.__file__).with_name("BOM-template.xlsx")).active
            self.assertAlmostEqual(sheet.column_dimensions["E"].width,
                                   template.column_dimensions["E"].width * 1.20)
            self.assertEqual(sheet.column_dimensions["F"].width, 34.0)
            self.assertIsNone(sheet.row_dimensions[9].height)
            self.assertFalse(sheet.row_dimensions[9].customHeight)
            self.assertTrue(all(sheet.cell(9, c).alignment.wrap_text for c in range(1, 7)))
            self.assertEqual(sheet.row_dimensions[8].height, 30)
            self.assertEqual(sheet.print_title_rows, "$1:$8")
            self.assertIn("$A$1:$F$9", sheet.print_area)
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 0)

    def test_numbered_worksheets_only_replace_counters(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            library = root / "library"
            sheets = library / "drawing-sheets"
            sheets.mkdir(parents=True)
            assembly_source = sheets / "alex-generic-pcba.kicad_wks"
            schematic_source = sheets / "alex-generic-sch.kicad_wks"
            assembly_original = b'\xef\xbb\xbf(kicad_wks (tbtext "PCBA ${#}/${##} ${VARIANT} ${ProjectTitle}"))\r\n'
            schematic_original = b'\xef\xbb\xbf(kicad_wks (tbtext "SCH ${#}/${##} ${VARIANT} ${ProjectTitle}"))\r\n'
            assembly_source.write_bytes(assembly_original)
            schematic_source.write_bytes(schematic_original)
            project = root / "project"
            project.mkdir()
            work = root / "output"
            work.mkdir()
            with (patch.dict(os.environ, KICAD_LIB_ROOT=str(library),
                             JOBSET_OUTPUT_WORK_PATH=str(work)),
                  patch.object(formatter, "assembly_page_count", return_value=2)):
                with patch.object(formatter, "assembly_variant_name", return_value="No Variant"):
                    formatter.prepare_assembly_worksheets(project)
                self.assertEqual((work / "_work/schematic.kicad_wks").read_bytes(),
                                 schematic_original.replace(b"${##}", b"1").replace(b"${#}", b"1")
                                 .replace(b"${VARIANT}", b"No Variant"))
                for page, view in ((1, "top"), (2, "bottom")):
                    self.assertEqual((work / f"_work/assembly-{view}.kicad_wks").read_bytes(),
                                     assembly_original.replace(b"${##}", b"2")
                                     .replace(b"${#}", str(page).encode())
                                     .replace(b"${VARIANT}", b"No Variant"))
                self.assertEqual(assembly_source.read_bytes(), assembly_original)
                self.assertEqual(schematic_source.read_bytes(), schematic_original)
                # A stale or redirected target cannot be silently overwritten.
                with self.assertRaises(FileExistsError):
                    with patch.object(formatter, "assembly_variant_name", return_value="No Variant"):
                        formatter.prepare_assembly_worksheets(project)

    def test_top_only_worksheet_uses_one_of_one(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sheets = root / "library/drawing-sheets"
            sheets.mkdir(parents=True)
            (sheets / "alex-generic-pcba.kicad_wks").write_bytes(
                b'(tbtext "PCBA ${#}/${##} ${VARIANT}")')
            (sheets / "alex-generic-sch.kicad_wks").write_bytes(
                b'(tbtext "SCH ${#}/${##} ${VARIANT}")')
            project = root / "project"
            project.mkdir()
            output = root / "output"
            output.mkdir()
            with (patch.dict(os.environ, KICAD_LIB_ROOT=str(root / "library"),
                             JOBSET_OUTPUT_WORK_PATH=str(output)),
                  patch.object(formatter, "assembly_page_count", return_value=1)):
                with patch.object(formatter, "assembly_variant_name", return_value="No Variant"):
                    formatter.prepare_assembly_worksheets(project)
            self.assertEqual((output / "_work/assembly-top.kicad_wks").read_bytes(),
                             b'(tbtext "PCBA 1/1 No Variant")')
            self.assertEqual((output / "_work/schematic.kicad_wks").read_bytes(),
                             b'(tbtext "SCH 1/1 No Variant")')
            self.assertFalse((output / "_work/assembly-bottom.kicad_wks").exists())

    def test_fabrication_worksheet_is_temporary_exact_copy(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sheets = root / "library/drawing-sheets"
            sheets.mkdir(parents=True)
            source = sheets / "alex-generic-pcb.kicad_wks"
            original = b'\xef\xbb\xbf(kicad_wks (tbtext "${ProjectTitle}"))\r\n'
            source.write_bytes(original)
            project = root / "project"
            project.mkdir()
            output = root / "output"
            output.mkdir()
            with patch.dict(os.environ, KICAD_LIB_ROOT=str(root / "library"),
                            JOBSET_OUTPUT_WORK_PATH=str(output)):
                formatter.prepare_fabrication_worksheet(project)
            self.assertEqual((output / "_work/fabrication.kicad_wks").read_bytes(), original)
            self.assertEqual(source.read_bytes(), original)

    def test_missing_canonical_worksheet_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            project = root / "project"
            project.mkdir()
            output = root / "output"
            output.mkdir()
            with (patch.dict(os.environ, KICAD_LIB_ROOT=directory,
                             JOBSET_OUTPUT_WORK_PATH=str(output)),
                  patch.object(formatter, "assembly_page_count", return_value=1),
                  patch.object(formatter, "assembly_variant_name", return_value="No Variant")):
                with self.assertRaisesRegex(FileNotFoundError, "KICAD_LIB_ROOT"):
                    formatter.prepare_assembly_worksheets(project)

    def test_rejects_unexpected_bom_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "bom.csv"
            source.write_text("wrong,columns\n1,2\n")
            with self.assertRaisesRegex(ValueError, "Unexpected native BOM columns"):
                formatter.create_bom(source, Path(directory) / "out.xlsx", {}, "", datetime.now())

    def test_merge_preserves_page_order_and_dimensions(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name, width in (("top", 100), ("bottom", 200)):
                writer = PdfWriter()
                writer.add_blank_page(width=width, height=300)
                writer.write(root / f"{name}.pdf")
            target = root / "assembly.pdf"
            formatter.merge_assembly(root / "top.pdf", root / "bottom.pdf", target, "Assembly")
            reader = PdfReader(target)
            self.assertEqual([page.mediabox.width for page in reader.pages], [100, 200])
            self.assertEqual([item.title for item in reader.outline], ["Top", "Bottom"])

    def test_top_only_assembly_is_published_without_rewriting(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=300)
            writer.add_metadata({"/Title": "Native top drawing"})
            top = root / "top.pdf"
            writer.write(top)
            target = root / "assembly.pdf"
            formatter.merge_assembly(top, None, target, "Ignored merge title")
            self.assertEqual(target.read_bytes(), top.read_bytes())
            reader = PdfReader(target)
            self.assertEqual(len(reader.pages), 1)
            self.assertEqual(reader.metadata.title, "Native top drawing")


if __name__ == "__main__":
    unittest.main()
