"""Format a KiCad project's outputs inside a Prism-selected Jobset destination.

This does not run exports, create releases, invoke Git, or publish anything.
KiCad supplies the project root explicitly and the temporary destination through
JOBSET_OUTPUT_WORK_PATH.
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime
import json
import os
from pathlib import Path
import re
import shutil
import struct
import sys
from tempfile import TemporaryDirectory


HEADERS = [
    "Line #", "Quantity", "Designator", "Description", "Manufacturer",
    "Manufacturer Part Number",
]


def filename_value(value: str) -> str:
    if not value or re.search(r'[<>:"/\\|?*\x00-\x1f]', value) or value.endswith((" ", ".")):
        raise ValueError(f"Not a portable filename value: {value!r}")
    return value


def create_bom(source: Path, target: Path, variables: dict, variant: str, published: datetime) -> None:
    from copy import copy
    import openpyxl
    from openpyxl.styles import Alignment, Font

    with source.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != HEADERS:
            raise ValueError(f"Unexpected native BOM columns: {reader.fieldnames}")
        rows = list(reader)

    workbook = openpyxl.load_workbook(Path(__file__).with_name("BOM-template.xlsx"))
    sheet = workbook.active
    manufacturer = sheet.column_dimensions["E"]
    baseline = manufacturer.width
    if baseline is None:
        baseline = sheet.sheet_format.defaultColWidth
    manufacturer.width = (baseline if baseline is not None else 13.0) * 1.20
    # Calibri 11 at width 34 fits representative long orderable numbers (roughly
    # 27 characters) with padding; wrap remains enabled for uncommon outliers.
    sheet.column_dimensions["F"].width = 34.0
    if [sheet.cell(8, c).value for c in range(1, 7)] != HEADERS:
        raise ValueError("BOM template header changed; update field mapping before release")

    sheet["C2"] = variables["ProjectTitle"]
    sheet["C3"] = variables.get("ProjectPartNumber") or "N/A"
    sheet["C4"] = str(variables["ProjectPCBARevision"])
    sheet["C5"] = published.replace(tzinfo=None)
    sheet["C5"].number_format = "yyyy-mm-dd hh-mm"
    sheet["C6"] = variant
    # Revisions, identifiers and descriptions are text, never spreadsheet formulas.
    for address in ("C2", "C3", "C4", "C6"):
        sheet[address].data_type = "s"
        sheet[address].number_format = "@"

    for row_number, row in enumerate(rows, 9):
        for column, header in enumerate(HEADERS, 1):
            cell = sheet.cell(row_number, column)
            value = row[header]
            cell.value = int(value) if column <= 2 else value
            if column > 2:
                cell.data_type = "s"
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if column <= 2 else "left")
            cell.border = copy(sheet.cell(8, column).border)
        # Leave wrapped item rows eligible for the spreadsheet reader's AutoFit.
        sheet.row_dimensions[row_number].height = None

    for col in range(1, 7):
        alignment = copy(sheet.cell(8, col).alignment)
        alignment.wrap_text = True
        sheet.cell(8, col).alignment = alignment
    sheet.row_dimensions[8].height = max(sheet.row_dimensions[8].height or 0, 30)
    sheet.print_area = f"A1:F{max(8, 8 + len(rows))}"
    sheet.print_title_rows = "1:8"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    workbook.save(target)


def merge_assembly(top: Path, bottom: Path | None, target: Path, title: str) -> None:
    from pypdf import PdfReader, PdfWriter

    if bottom is None:
        reader = PdfReader(top)
        if len(reader.pages) != 1:
            raise ValueError(f"Expected one Top drawing page, got {len(reader.pages)}")
        shutil.copyfile(top, target)
        return

    writer = PdfWriter()
    for source, label in ((top, "Top"), (bottom, "Bottom")):
        reader = PdfReader(source)
        if len(reader.pages) != 1:
            raise ValueError(f"Expected one {label} drawing page, got {len(reader.pages)}")
        writer.append(reader, outline_item=label, import_outline=False)
    writer.add_metadata({"/Title": title, "/Subject": "PCBA assembly: top and mirrored bottom"})
    with target.open("wb") as stream:
        writer.write(stream)


def recolor_u3d_pads(data: bytes) -> bytes:
    """Give KiCad's named pad material the same diffuse RGB as its copper material."""
    material_block = 0xFFFFFF54
    materials: dict[str, tuple[int, bytes]] = {}
    offset = 0

    while offset + 12 <= len(data):
        block_type, data_size, metadata_size = struct.unpack_from("<III", data, offset)
        payload = offset + 12
        data_end = payload + data_size
        block_end = (data_end + 3) & ~3
        block_end += (metadata_size + 3) & ~3
        if data_end > len(data) or block_end > len(data):
            raise ValueError("Malformed U3D block length")

        if block_type == material_block:
            if data_size < 2:
                raise ValueError("Malformed U3D material block")
            name_size = struct.unpack_from("<H", data, payload)[0]
            name_end = payload + 2 + name_size
            # Attributes, ambient RGB, then diffuse RGB.
            diffuse = name_end + 4 + 12
            if diffuse + 12 > data_end:
                raise ValueError("Malformed U3D material values")
            try:
                name = data[payload + 2:name_end].decode("utf-8")
            except UnicodeDecodeError as error:
                raise ValueError("Invalid U3D material name") from error
            if name in materials:
                raise ValueError(f"Duplicate U3D material: {name}")
            materials[name] = (diffuse, data[diffuse:diffuse + 12])

        offset = block_end

    if offset != len(data):
        raise ValueError("Trailing data after final U3D block")
    missing = {"m_Copper_0", "m_Pads_0"} - materials.keys()
    if missing:
        raise ValueError(f"Required U3D material missing: {', '.join(sorted(missing))}")

    pad_offset, pad_rgb = materials["m_Pads_0"]
    _, copper_rgb = materials["m_Copper_0"]
    if pad_rgb == copper_rgb:
        return data
    result = bytearray(data)
    result[pad_offset:pad_offset + 12] = copper_rgb
    return bytes(result)


_STEP_ENTITY = re.compile(r"(?ms)^(#\d+)\s*=\s*(.*?);(?=\r?$)")
_STEP_REFERENCE = re.compile(r"#\d+")
_STEP_COLOUR = re.compile(
    r"^COLOUR_RGB\(('[^']*'),\s*([^,]+),\s*([^,]+),\s*([^\)]+)\)$", re.S)


def _step_entities(data: str) -> dict[str, str]:
    entities = {match.group(1): match.group(2) for match in _STEP_ENTITY.finditer(data)}
    if not entities:
        raise ValueError("STEP file contains no entities")
    return entities


def _step_unique_referrer(entities: dict[str, str], entity_type: str, target: str) -> str:
    matches = [identifier for identifier, body in entities.items()
               if body.startswith(entity_type + "(") and target in _STEP_REFERENCE.findall(body)]
    if len(matches) != 1:
        raise ValueError(f"Expected one STEP {entity_type} referencing {target}, got {len(matches)}")
    return matches[0]


def _step_product_colour(entities: dict[str, str], suffix: str) -> str:
    products = []
    for identifier, body in entities.items():
        match = re.match(r"^PRODUCT\('([^']*)'", body)
        if match and match.group(1).endswith(suffix):
            products.append(identifier)
    if len(products) != 1:
        raise ValueError(f"Expected one STEP product ending {suffix!r}, got {len(products)}")

    formation = _step_unique_referrer(entities, "PRODUCT_DEFINITION_FORMATION", products[0])
    definition = _step_unique_referrer(entities, "PRODUCT_DEFINITION", formation)
    definition_shape = _step_unique_referrer(entities, "PRODUCT_DEFINITION_SHAPE", definition)
    shape_link = _step_unique_referrer(entities, "SHAPE_DEFINITION_REPRESENTATION", definition_shape)
    link_refs = _STEP_REFERENCE.findall(entities[shape_link])
    if len(link_refs) != 2:
        raise ValueError(f"Unexpected STEP shape link for product ending {suffix!r}")
    representation = link_refs[1]
    representation_refs = _STEP_REFERENCE.findall(entities[representation])
    if not representation_refs:
        raise ValueError(f"STEP representation has no context for product ending {suffix!r}")
    context = representation_refs[-1]

    presentations = [identifier for identifier, body in entities.items()
                     if body.startswith("MECHANICAL_DESIGN_GEOMETRIC_PRESENTATION_REPRESENTATION(")
                     and _STEP_REFERENCE.findall(body)[-1:] == [context]]
    if len(presentations) != 1:
        raise ValueError(f"Expected one STEP presentation for product ending {suffix!r}")

    colours: set[str] = set()
    for styled in _STEP_REFERENCE.findall(entities[presentations[0]])[:-1]:
        styled_refs = _STEP_REFERENCE.findall(entities.get(styled, ""))
        if not styled_refs:
            raise ValueError(f"Malformed STEP style for product ending {suffix!r}")
        pending = [styled_refs[0]]  # Ignore the styled geometry reference.
        visited: set[str] = set()
        while pending:
            identifier = pending.pop()
            if identifier in visited:
                continue
            visited.add(identifier)
            body = entities.get(identifier, "")
            if body.startswith("COLOUR_RGB("):
                colours.add(identifier)
            else:
                pending.extend(_STEP_REFERENCE.findall(body))
    if len(colours) != 1:
        raise ValueError(f"Expected one STEP colour for product ending {suffix!r}, got {len(colours)}")
    return next(iter(colours))


def recolor_step_pads(data: str) -> str:
    """Give KiCad's named pad STEP product the same RGB as its copper product."""
    entities = _step_entities(data)
    copper_id = _step_product_colour(entities, "_copper")
    pad_id = _step_product_colour(entities, "_pad")
    copper = _STEP_COLOUR.match(entities[copper_id])
    pad = _STEP_COLOUR.match(entities[pad_id])
    if copper is None or pad is None:
        raise ValueError("Malformed STEP copper or pad colour")
    if copper.groups()[1:] == pad.groups()[1:]:
        return data
    replacement = f"COLOUR_RGB({pad.group(1)},{copper.group(2)},{copper.group(3)},{copper.group(4)})"
    match = next(match for match in _STEP_ENTITY.finditer(data) if match.group(1) == pad_id)
    return data[:match.start(2)] + replacement + data[match.end(2):]


def recolor_step_file(source: Path, target: Path) -> None:
    with source.open(encoding="utf-8", newline="") as stream:
        data = stream.read()
    changed = recolor_step_pads(data)
    with target.open("w", encoding="utf-8", newline="") as stream:
        stream.write(changed)
    with target.open(encoding="utf-8", newline="") as stream:
        checked = stream.read()
    if recolor_step_pads(checked) != checked:
        raise ValueError("Pad STEP colour did not persist")


def recolor_3d_pdf(source: Path, target: Path) -> None:
    """Recolor the pad material in KiCad's single-page 3D PDF output."""
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(source)
    if len(reader.pages) != 1:
        raise ValueError(f"Expected one 3D PDF page, got {len(reader.pages)}")
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    annotations = writer.pages[0].get("/Annots", [])
    three_d = [item.get_object() for item in annotations
               if item.get_object().get("/Subtype") == "/3D"]
    if len(three_d) != 1:
        raise ValueError(f"Expected one 3D annotation, got {len(three_d)}")
    stream = three_d[0].get("/3DD")
    if stream is None:
        raise ValueError("3D annotation has no embedded stream")
    stream = stream.get_object()
    if stream.get("/Subtype") != "/U3D":
        raise ValueError(f"Expected embedded U3D, got {stream.get('/Subtype')}")
    stream.set_data(recolor_u3d_pads(stream.get_data()))
    with target.open("wb") as output:
        writer.write(output)

    check = PdfReader(target)
    annotation = check.pages[0]["/Annots"][0].get_object()
    checked_stream = annotation["/3DD"].get_object()
    checked = checked_stream.get_data()
    if recolor_u3d_pads(checked) != checked:
        raise ValueError("Pad material recolor did not persist")


def jobset_work_root(project_root: Path) -> Path:
    temporary = os.environ.get("JOBSET_OUTPUT_WORK_PATH")
    if not temporary:
        raise RuntimeError("Run this formatter only through Outputs.kicad_jobset in KiCad/Prism")
    root = Path(temporary).resolve(strict=True)
    if root == project_root or root in project_root.parents or project_root in root.parents:
        raise RuntimeError("Formatter output must be KiCad's temporary destination, outside source")
    return root


def clean_release_destination(project_root: Path, prepare_kind: str) -> Path:
    """Remove only the project output directory selected by this preparation job."""
    project_root = project_root.resolve(strict=True)
    jobset = json.loads((project_root / "Outputs.kicad_jobset").read_text(encoding="utf-8"))
    marker = f"--kind {prepare_kind}"
    prepare_jobs = [job["id"] for job in jobset["jobs"]
                    if marker in job.get("settings", {}).get("command", "")]
    if len(prepare_jobs) != 1:
        raise ValueError(f"Expected one {prepare_kind} job")
    destinations = [output for output in jobset["outputs"]
                    if prepare_jobs[0] in output.get("only", [])]
    if len(destinations) != 1:
        raise ValueError(f"Expected {prepare_kind} in exactly one Jobset destination")
    output_path = Path(destinations[0].get("settings", {}).get("output_path", ""))
    if (not output_path.parts or output_path.is_absolute()
            or any(part in ("", ".", "..") for part in output_path.parts)):
        raise ValueError(f"Jobset destination must be a relative project path: {output_path}")

    destination = project_root.joinpath(output_path)
    current = project_root
    for part in output_path.parts:
        current = current / part
        if current.is_symlink():
            raise ValueError(f"Refusing symlinked release destination path: {current}")
    resolved = destination.resolve()
    if resolved == project_root or project_root not in resolved.parents:
        raise ValueError(f"Jobset destination escapes the project: {resolved}")

    if destination.exists():
        if not destination.is_dir():
            raise ValueError(f"Jobset destination is not a directory: {destination}")
        for parent, directories, files in os.walk(destination, followlinks=False):
            for name in [*directories, *files]:
                item = Path(parent, name)
                if item.is_symlink():
                    raise ValueError(f"Refusing release destination containing symlink: {item}")
        shutil.rmtree(destination)
        print(f"Removed previous release destination {output_path.as_posix()}")
    return destination


def assembly_page_count(jobset_path: Path) -> int:
    jobset = json.loads(jobset_path.read_text(encoding="utf-8"))
    bottom_jobs = [job["id"] for job in jobset["jobs"]
                   if job.get("settings", {}).get("output_filename") == "_work/assembly-bottom.pdf"]
    prepare_jobs = [job["id"] for job in jobset["jobs"]
                    if "--kind prepare-assembly" in job.get("settings", {}).get("command", "")]
    if len(bottom_jobs) != 1 or len(prepare_jobs) != 1:
        raise ValueError("Expected one Bottom drawing job and one worksheet preparation job")
    destinations = [output for output in jobset["outputs"]
                    if prepare_jobs[0] in output.get("only", [])]
    if len(destinations) != 1:
        raise ValueError("Expected worksheet preparation in exactly one Jobset destination")
    return 2 if bottom_jobs[0] in destinations[0]["only"] else 1


def assembly_variant_name(jobset_path: Path) -> str:
    jobset = json.loads(jobset_path.read_text(encoding="utf-8"))
    prepare_jobs = [job["id"] for job in jobset["jobs"]
                    if "--kind prepare-assembly" in job.get("settings", {}).get("command", "")]
    if len(prepare_jobs) != 1:
        raise ValueError("Expected one worksheet preparation job")
    destinations = [output for output in jobset["outputs"]
                    if prepare_jobs[0] in output.get("only", [])]
    if len(destinations) != 1:
        raise ValueError("Expected worksheet preparation in exactly one Jobset destination")
    enabled = set(destinations[0]["only"])
    variants = set()
    for job in jobset["jobs"]:
        if job["id"] not in enabled:
            continue
        settings = job.get("settings", {})
        for key in ("variant", "variant_name"):
            value = settings.get(key, "")
            if value:
                variants.add(value)
    if len(variants) > 1:
        raise ValueError(f"Assembly destination selects conflicting variants: {sorted(variants)}")
    return next(iter(variants), "No Variant")


def library_root() -> Path:
    return Path(os.environ.get("KICAD_LIB_ROOT") or Path(__file__).resolve().parent.parent)


def write_temporary_worksheet(target: Path, content: bytes) -> None:
    # Exclusive creation prevents following or overwriting a pre-existing target/symlink.
    with target.open("xb") as stream:
        stream.write(content)


def prepare_fabrication_worksheet(project_root: Path) -> None:
    root = jobset_work_root(project_root)
    clean_release_destination(project_root, "prepare-fabrication")
    source = library_root() / "drawing-sheets" / "alex-generic-pcb.kicad_wks"
    if not source.is_file():
        raise FileNotFoundError(f"Canonical PCB worksheet missing: {source}. "
                                "Set KICAD_LIB_ROOT to the shared kicad-lib checkout.")
    content = source.read_bytes()
    if any(token not in content for token in (b"${#}", b"${##}")):
        raise ValueError(f"Canonical fabrication worksheet must contain page tokens: {source}")
    content = content.replace(b"${##}", b"1").replace(b"${#}", b"1")
    work = root / "_work"
    work.mkdir(exist_ok=True)
    write_temporary_worksheet(work / "fabrication.kicad_wks", content)
    print(f"Prepared fabrication worksheet from {source.resolve()}")


def prepare_assembly_worksheets(project_root: Path) -> None:
    root = jobset_work_root(project_root)
    clean_release_destination(project_root, "prepare-assembly")
    jobset = project_root / "Outputs.kicad_jobset"
    page_count = assembly_page_count(jobset)
    variant = assembly_variant_name(jobset).encode()
    sheets = library_root() / "drawing-sheets"
    schematic_source = sheets / "alex-generic-sch.kicad_wks"
    assembly_source = sheets / "alex-generic-pcba.kicad_wks"
    for source in (schematic_source, assembly_source):
        if not source.is_file():
            raise FileNotFoundError(f"Canonical worksheet missing: {source}. "
                                    "Set KICAD_LIB_ROOT to the shared kicad-lib checkout.")
    # Bytes preserve encoding, BOM and line endings; all other variables stay native.
    schematic = schematic_source.read_bytes()
    assembly = assembly_source.read_bytes()
    for source, content in ((schematic_source, schematic), (assembly_source, assembly)):
        if any(token not in content for token in (b"${#}", b"${##}", b"${VARIANT}")):
            raise ValueError(f"Canonical worksheet must contain page and variant tokens: {source}")
    work = root / "_work"
    work.mkdir(exist_ok=True)
    write_temporary_worksheet(work / "schematic.kicad_wks",
                              schematic.replace(b"${##}", b"1").replace(b"${#}", b"1")
                              .replace(b"${VARIANT}", variant))
    views = ((1, "top"), (2, "bottom")) if page_count == 2 else ((1, "top"),)
    for page, view in views:
        target = work / f"assembly-{view}.kicad_wks"
        write_temporary_worksheet(target,
                                  assembly.replace(b"${##}", str(page_count).encode())
                                  .replace(b"${#}", str(page).encode())
                                  .replace(b"${VARIANT}", variant))
    print(f"Prepared schematic and assembly worksheets for variant {variant.decode()}")


def finalize(kind: str, project_root: Path, variant_name: str = "") -> list[Path]:
    root = jobset_work_root(project_root)
    projects = list(project_root.glob("*.kicad_pro"))
    if len(projects) != 1:
        raise RuntimeError("Expected exactly one KiCad project alongside the canonical Jobset")
    variables = json.loads(projects[0].read_text(encoding="utf-8"))["text_variables"]
    title = filename_value(variables["ProjectTitle"])
    revision = filename_value(str(variables["ProjectPCBRevision" if kind == "fabrication" else "ProjectPCBARevision"]))
    variant = filename_value(variant_name or "No Variant")
    published = datetime.now().astimezone().replace(second=0, microsecond=0)
    suffix = f" - {variant} ({published:%Y-%m-%d %H-%M})"
    prefix = f"{title} R{revision} "
    work = root / "_work"

    if kind == "fabrication":
        mapping = {"fabrication.pdf": ("FAB", ".pdf"), "envelope.step": ("STEP", ".step")}
        inputs = list(mapping)
        destination = root
    else:
        mapping = {"schematic.pdf": ("SCH", ".pdf"), "assembly-3d.pdf": ("3DPCB", ".pdf"),
                   "positions-all-pos.csv": ("Pick Place", ".csv")}
        inputs = list(mapping) + ["assembly-top.pdf", "bom.csv"]
        destination = root / variant

    for name in inputs:
        source = work / name
        if source.is_symlink() or not source.is_file() or not source.stat().st_size:
            raise RuntimeError(f"Missing or invalid native output: {name}")

    assembly_bottom = None
    if kind == "assembly":
        candidate = work / "assembly-bottom.pdf"
        if candidate.is_symlink() or (candidate.exists()
                                      and (not candidate.is_file() or not candidate.stat().st_size)):
            raise RuntimeError("Missing or invalid native output: assembly-bottom.pdf")
        if candidate.is_file():
            assembly_bottom = candidate
            inputs.append(candidate.name)

    # Build every formatted file before placing any of them in the destination.
    with TemporaryDirectory(prefix="format-", dir=root) as staging_path:
        staging = Path(staging_path)
        for name, (file_type, extension) in mapping.items():
            target = staging / (prefix + file_type + suffix + extension)
            if name == "assembly-3d.pdf":
                recolor_3d_pdf(work / name, target)
            elif name == "envelope.step":
                recolor_step_file(work / name, target)
            else:
                shutil.copyfile(work / name, target)
        if kind == "assembly":
            merge_assembly(work / "assembly-top.pdf", assembly_bottom,
                           staging / (prefix + "ASY" + suffix + ".pdf"), prefix + "ASY" + suffix)
            create_bom(work / "bom.csv", staging / (prefix + "BOM" + suffix + ".xlsx"),
                       variables, variant, published)
        destination.mkdir(parents=True, exist_ok=True)
        result = []
        for source in sorted(staging.iterdir()):
            target = destination / source.name
            source.replace(target)
            result.append(target)

    for name in inputs:
        (work / name).unlink()
    if kind == "fabrication":
        (work / "fabrication.kicad_wks").unlink(missing_ok=True)
    else:
        (work / "schematic.kicad_wks").unlink(missing_ok=True)
        for view in ("top", "bottom"):
            (work / f"assembly-{view}.kicad_wks").unlink(missing_ok=True)
    work.rmdir()
    for path in result:
        print(f"Created {path.relative_to(root).as_posix()}")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", required=True, type=Path,
                        help="KiCad project directory supplied by ${KIPRJMOD}")
    parser.add_argument("--kind", choices=("fabrication", "assembly", "prepare-fabrication",
                                           "prepare-assembly"), required=True)
    parser.add_argument("--variant", default="", help="Empty selects the base design, labelled No Variant")
    args = parser.parse_args()
    try:
        project_root = args.project_root.resolve(strict=True)
        if not project_root.is_dir():
            raise NotADirectoryError(f"Project root is not a directory: {project_root}")
        if args.kind == "prepare-fabrication":
            prepare_fabrication_worksheet(project_root)
        elif args.kind == "prepare-assembly":
            prepare_assembly_worksheets(project_root)
        else:
            finalize(args.kind, project_root, args.variant)
    except Exception as error:
        print(f"Output formatting failed: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
